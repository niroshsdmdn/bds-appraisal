"""Turn an uploaded proposal (PDF, XLSX/XLSM or CSV) into one common shape.

Every format becomes:
  * grid  - table rows in document order (list of cell strings)
  * lines - readable text lines in document order
  * images - embedded pictures (beneficiary / site photos)

The parser in parse.py only ever sees this structure, so all three formats
are assessed by the same rules.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field

SUPPORTED = ("pdf", "xlsx", "xlsm", "csv")

# pdf text layers from the BDS template often split numbers: "1 0,000.00"
_SPLIT_BEFORE_GROUP = re.compile(r"(?<=\d) (?=\d{0,2},\d{3}\b)")
_SPLIT_BEFORE_COMMA = re.compile(r"(?<=\d) (?=,\d{3}\b)")
_NUM_CELL = re.compile(r"^-?\d+(?:\.\d+)?$")


@dataclass
class ImageInfo:
    page: int
    width: float
    height: float
    data: bytes | None = None
    media_type: str = "image/png"


@dataclass
class Extracted:
    file_name: str
    file_type: str
    raw_bytes: bytes
    grid: list[list[str]] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)
    images: list[ImageInfo] = field(default_factory=list)
    page_count: int = 0
    sheet_names: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n".join(self.lines)

    @property
    def sha1(self) -> str:
        return hashlib.sha1(self.raw_bytes).hexdigest()[:12]

    @property
    def photo_count(self) -> int:
        # logos and icons are small; photos are not
        return sum(1 for im in self.images if im.width >= 90 and im.height >= 90)

    def grid_as_text(self) -> str:
        out = []
        for row in self.grid:
            cells = [c for c in row if c]
            if cells:
                out.append(" | ".join(cells))
        return "\n".join(out)


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------
def fix_split_numbers(text: str) -> str:
    text = _SPLIT_BEFORE_GROUP.sub("", text)
    return _SPLIT_BEFORE_COMMA.sub("", text)


def to_number(cell) -> float | None:
    """Parse a single table cell. '-' counts as zero (template convention)."""
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        return float(cell)
    s = str(cell).strip()
    if s in {"-", "–", "—"}:
        return 0.0
    s = re.sub(r"(?i)^rs\.?", "", s)
    s = s.replace(",", "").replace(" ", "").rstrip("%")
    if _NUM_CELL.match(s):
        return float(s)
    return None


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.2f}"
    return str(v).strip()


# --------------------------------------------------------------------------
# format readers
# --------------------------------------------------------------------------
def _read_pdf(ex: Extracted) -> None:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(ex.raw_bytes)) as pdf:
        ex.page_count = len(pdf.pages)
        text_chars = 0
        for pno, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            text_chars += len(text.strip())
            for line in text.splitlines():
                line = fix_split_numbers(line.strip())
                if line:
                    ex.lines.append(line)
            for table in page.extract_tables():
                for row in table:
                    ex.grid.append([(c or "").strip() for c in row])
            for im in page.images:
                data = None
                try:
                    bbox = (max(im["x0"], 0), max(im["top"], 0),
                            min(im["x1"], page.width), min(im["bottom"], page.height))
                    buf = io.BytesIO()
                    page.crop(bbox).to_image(resolution=110).save(buf, format="PNG")
                    data = buf.getvalue()
                except Exception:  # image rendering is best-effort
                    pass
                ex.images.append(ImageInfo(pno, float(im["width"]), float(im["height"]), data))
    if text_chars < 200:
        ex.warnings.append(
            "Very little machine-readable text was found. The PDF may be a scan; "
            "turn on the AI review so the pages are read visually."
        )


def _read_xlsx(ex: Extracted) -> None:
    from openpyxl import load_workbook

    wb_values = load_workbook(io.BytesIO(ex.raw_bytes), data_only=True)
    wb_formulas = load_workbook(io.BytesIO(ex.raw_bytes), data_only=False)
    missing_cache = 0
    for ws in wb_values.worksheets:
        if ws.sheet_state != "visible":
            continue
        ex.sheet_names.append(ws.title)
        wf = wb_formulas[ws.title]
        ex.lines.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = c.value
                if v is None:
                    f = wf[c.coordinate].value
                    if isinstance(f, str) and f.startswith("="):
                        missing_cache += 1
                cells.append(_cell_to_str(v))
            if any(cells):
                ex.grid.append(cells)
                ex.lines.append(" ".join(c for c in cells if c))
        for img in getattr(ws, "_images", []):
            try:
                data = img._data()
            except Exception:
                data = None
            ex.images.append(ImageInfo(0, float(img.width or 0), float(img.height or 0), data,
                                       "image/png" if (data or b"")[:4] == b"\x89PNG" else "image/jpeg"))
    ex.page_count = len(ex.sheet_names)
    if missing_cache:
        ex.warnings.append(
            f"{missing_cache} formula cells have no saved result. Open the workbook in "
            "Excel, let it recalculate, save, and upload again for complete figures."
        )


def _read_csv(ex: Extracted) -> None:
    raw = ex.raw_bytes
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    for row in csv.reader(io.StringIO(text), dialect):
        cells = [c.strip() for c in row]
        if any(cells):
            ex.grid.append(cells)
            ex.lines.append(" ".join(c for c in cells if c))
    ex.page_count = 1


def extract(file_name: str, data: bytes) -> Extracted:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '.{ext}'. Upload a PDF, XLSX or CSV file.")
    ex = Extracted(file_name=file_name, file_type="xlsx" if ext == "xlsm" else ext, raw_bytes=data)
    {"pdf": _read_pdf, "xlsx": _read_xlsx, "csv": _read_csv}[ex.file_type](ex)
    if not ex.grid and not ex.lines:
        ex.warnings.append("No content could be read from this file.")
    return ex
