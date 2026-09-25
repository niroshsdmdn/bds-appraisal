"""Appraisal report exports: an Excel workbook and a printable HTML page."""
from __future__ import annotations

import html
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .checks import Appraisal
from .config import BANDS
from .parse import Proposal

INK = "1C2A3A"
NAVY = "1F3A68"
TINTS = {"pass": "E3EFE5", "partial": "FBF0DA", "fail": "F6E0DC", "na": "EEEEEE", "info": "E8EEF6"}
FONT = "Arial"


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autowidth(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _body_font(ws) -> None:
    thin = Side(style="thin", color="D5DAE1")
    for row in ws.iter_rows():
        for cell in row:
            if cell.font is None or not cell.font.bold:
                cell.font = Font(name=FONT, color=INK, bold=cell.font.bold if cell.font else False)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value is not None:
                cell.border = Border(bottom=thin)


def build_excel(p: Proposal, a: Appraisal, meta: dict) -> bytes:
    wb = Workbook()
    dec = a.decision()

    # --- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "BDS livelihood proposal appraisal"
    ws["A1"].font = Font(name=FONT, size=15, bold=True, color=NAVY)
    rows = [
        ("File", meta.get("file_name", "")),
        ("Appraised on", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Appraised by", f"{meta.get('user', '')} ({meta.get('role', '')})"),
        ("Sector used for the checks", a.context.get("sector_name", "")),
        ("Total mark (out of 100)", a.total),
        ("System recommendation", dec["label"]),
        ("Guidance", dec["guidance"]),
        ("Critical findings", ", ".join(c.id for c in a.critical) or "None"),
        ("AI review", meta.get("ai_model") or "Not used"),
        ("Reviewer recommendation", meta.get("reviewer_decision", "")),
        ("Reviewer comments", meta.get("reviewer_comment", "")),
        ("Approver decision", meta.get("approver_decision", "")),
        ("Approver comments", meta.get("approver_comment", "")),
        ("Conditions", "\n".join(meta.get("conditions", []))),
    ]
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, color=INK)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Decision bands").font = Font(name=FONT, bold=True, color=NAVY)
    for minimum, label, _, guidance in BANDS:
        r += 1
        ws.cell(row=r, column=1, value=f"{minimum}+")
        ws.cell(row=r, column=2, value=f"{label}: {guidance}")
    r += 2
    ws.cell(row=r, column=1, value="Proposal at a glance").font = Font(name=FONT, bold=True, color=NAVY)
    for k, v in p.as_summary_rows():
        r += 1
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    _body_font(ws)
    ws["A1"].font = Font(name=FONT, size=15, bold=True, color=NAVY)
    _autowidth(ws, {1: 34, 2: 90})

    # --- Criteria
    ws = wb.create_sheet("Criteria")
    head = ["ID", "Criterion", "Max", "Rule score", "AI score", "Final score", "Reviewer note"]
    ws.append(head)
    for cs in a.criteria.values():
        ws.append([cs.id, cs.name, cs.max, cs.rule_score, cs.ai_score, cs.score, cs.reviewer_note])
    n = len(a.criteria) + 1
    ws.append(["", "Total", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", "", f"=SUM(F2:F{n})", ""])
    for col in (3, 4, 5, 6):
        for row in range(2, n + 2):
            ws.cell(row=row, column=col).number_format = "0.0;-0.0;-"
    _body_font(ws)
    _style_header(ws, 1, len(head))
    for col in range(1, len(head) + 1):
        ws.cell(row=n + 1, column=col).font = Font(name=FONT, bold=True, color=INK)
    _autowidth(ws, {1: 6, 2: 42, 3: 8, 4: 11, 5: 10, 6: 11, 7: 50})

    # --- Findings
    ws = wb.create_sheet("Findings")
    head = ["Check", "Criterion", "Test", "Status", "Severity", "Marks", "Out of", "Finding", "Source"]
    ws.append(head)
    for c in a.checks:
        ws.append([c.id, c.criterion, c.title, c.status, c.severity, c.earned, c.weight, c.finding, c.source])
    _body_font(ws)
    _style_header(ws, 1, len(head))
    for i, c in enumerate(a.checks, start=2):
        ws.cell(row=i, column=4).fill = PatternFill("solid", fgColor=TINTS.get(c.status, "FFFFFF"))
    _autowidth(ws, {1: 8, 2: 9, 3: 36, 4: 9, 5: 10, 6: 8, 7: 8, 8: 90, 9: 8})
    ws.freeze_panes = "A2"

    # --- Improvement plan
    ws = wb.create_sheet("Improvement plan")
    head = ["Priority", "Check", "Criterion", "Severity", "Issue", "Action for the writer", "Marks recoverable", "Source"]
    ws.append(head)
    for i, row in enumerate(a.improvement_plan(), start=1):
        ws.append([i, row["check"], row["criterion_name"], row["severity"], row["issue"],
                   row["action"], row["mark_gain"], row["source"]])
    _body_font(ws)
    _style_header(ws, 1, len(head))
    _autowidth(ws, {1: 8, 2: 8, 3: 28, 4: 10, 5: 60, 6: 60, 7: 12, 8: 8})
    ws.freeze_panes = "A2"

    # --- Extracted data
    ws = wb.create_sheet("Extracted data")
    ws.append(["Section", "Item", "Value 1", "Value 2", "Value 3", "Value 4"])
    for m in p.members:
        ws.append(["Family", m.name, m.age, m.relationship, m.health])
    for label, b in (("Current budget", p.current), ("Expected budget", p.expected)):
        for inc in b.incomes:
            ws.append([label, f"Income: {inc.source}", inc.earner, inc.amount])
        for k, v in b.expenses.items():
            ws.append([label, f"Expense: {k}", v])
        ws.append([label, "Totals (income, expense, members, per capita)", b.total_income,
                   b.total_expense, b.members, b.per_capita])
    for k, v in p.investment.items():
        ws.append(["Investment", k, v.get("total"), v.get("own"), v.get("own_proposed"), v.get("grant")])
    for k, v in p.annual.items():
        ws.append(["Annual (current, year 1)", k, *v[:2]])
    for it in p.equipment:
        ws.append(["Equipment", it.name, it.qty, it.unit_price, it.total, it.bds])
    for risk, mit in p.risks:
        ws.append(["Risk", risk, mit])
    for t, info in p.trainings:
        ws.append(["Training", t, info])
    for g in p.goals:
        ws.append(["Goal", g])
    _body_font(ws)
    _style_header(ws, 1, 6)
    _autowidth(ws, {1: 24, 2: 44, 3: 18, 4: 18, 5: 18, 6: 18})

    # --- AI review
    ai = meta.get("ai")
    if ai:
        ws = wb.create_sheet("AI review")
        ws.append(["Model", ai.get("_model", "")])
        ws.append(["Summary", ai.get("summary", "")])
        ws.append(["Recommended decision", ai.get("recommended_decision", "")])
        ws.append(["Conditions", "\n".join(ai.get("conditions", []))])
        ws.append(["Photo observations", ai.get("photo_observations", "")])
        ws.append([])
        ws.append(["Criterion", "AI score", "Rationale"])
        for row in ai.get("criteria", []):
            ws.append([row.get("id"), row.get("score"), row.get("rationale")])
        ws.append([])
        ws.append(["Disagreement with rule", "Comment"])
        for d in ai.get("disagreements", []):
            ws.append([d.get("check"), d.get("comment")])
        _body_font(ws)
        _autowidth(ws, {1: 24, 2: 12, 3: 100})

    # --- Audit trail
    ws = wb.create_sheet("Audit trail")
    ws.append(["When", "Who", "Role", "Action", "Detail"])
    for e in meta.get("audit", []):
        ws.append([e.get("when"), e.get("who"), e.get("role"), e.get("action"), e.get("detail")])
    _body_font(ws)
    _style_header(ws, 1, 5)
    _autowidth(ws, {1: 20, 2: 20, 3: 12, 4: 24, 5: 70})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_html(p: Proposal, a: Appraisal, meta: dict) -> str:
    e = html.escape
    dec = a.decision()
    crit_rows = "".join(
        f"<tr><td>{e(c.id)}</td><td>{e(c.name)}</td><td class=n>{c.score:.1f}</td><td class=n>{c.max:g}</td></tr>"
        for c in a.criteria.values())
    plan_rows = "".join(
        f"<tr><td>{i}</td><td>{e(r['severity'])}</td><td>{e(r['issue'])}</td><td>{e(r['action'])}</td>"
        f"<td class=n>+{r['mark_gain']:.1f}</td></tr>"
        for i, r in enumerate(a.improvement_plan(), start=1))
    glance = "".join(f"<tr><th>{e(k)}</th><td>{e(str(v))}</td></tr>" for k, v in p.as_summary_rows())
    crit = ", ".join(f"{c.id} {c.finding}" for c in a.critical) or "None"
    return f"""<!doctype html><html><head><meta charset="utf-8">
<title>Appraisal - {e(p.name or meta.get('file_name', ''))}</title>
<style>
body{{font-family:'Source Serif 4',Georgia,serif;color:#1C2A3A;max-width:960px;margin:2rem auto;padding:0 1rem;line-height:1.5}}
h1{{font-size:1.6rem;margin:0 0 .25rem;color:#1F3A68}} h2{{font-size:1.15rem;margin-top:2rem;color:#1F3A68}}
table{{border-collapse:collapse;width:100%;font-size:.92rem}} td,th{{border-bottom:1px solid #D5DAE1;padding:.35rem .5rem;text-align:left;vertical-align:top}}
.n{{text-align:right;font-variant-numeric:tabular-nums}} .score{{font-size:2.4rem;font-weight:700}}
.dec{{display:inline-block;padding:.2rem .7rem;border-radius:4px;background:#1F3A68;color:#fff}}
@media print{{body{{margin:0}}}}
</style></head><body>
<h1>Livelihood proposal appraisal</h1>
<p>{e(meta.get('file_name', ''))} &middot; appraised {datetime.now():%d %b %Y} by {e(meta.get('user', ''))}</p>
<p><span class=score>{a.total:.1f}</span> / 100 &nbsp; <span class=dec>{e(dec['label'])}</span></p>
<p>{e(dec['guidance'])}</p>
<p><b>Critical findings:</b> {e(crit)}</p>
<p><b>Reviewer:</b> {e(meta.get('reviewer_decision', '') or '-')} {e(meta.get('reviewer_comment', ''))}<br>
<b>Approver:</b> {e(meta.get('approver_decision', '') or '-')} {e(meta.get('approver_comment', ''))}</p>
<h2>Marks by criterion</h2><table>{crit_rows}</table>
<h2>How to raise the mark</h2>
<table><tr><th>#</th><th>Severity</th><th>Issue</th><th>Action</th><th class=n>Marks</th></tr>{plan_rows}</table>
<h2>Proposal at a glance</h2><table>{glance}</table>
</body></html>"""
